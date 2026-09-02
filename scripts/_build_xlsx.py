"""Construit l'Excel campagnes + criteres + cases de modification pour Richard."""
from __future__ import annotations

import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
d = json.load(open(os.path.join(HERE, "data", "_research_dump.json"), encoding="utf-8"))
OUT = r"C:\Users\RI.GROS\Documents\CLAUDE\prospection-ia\CAMPAGNES_CRITERES.xlsx"

HEAD = PatternFill("solid", fgColor="1F4E78")
EDIT = PatternFill("solid", fgColor="FFF2CC")  # jaune = a remplir par Richard
ABMF = PatternFill("solid", fgColor="FCE4D6")
PERS = PatternFill("solid", fgColor="E2EFDA")
WHITEB = Font(color="FFFFFF", bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def kw_join(kws, n=12):
    used = [k["keyword"] for k in kws if k.get("used")]
    rest = [k["keyword"] for k in kws if not k.get("used")]
    allk = used + rest
    s = " · ".join(allk[:n])
    if len(allk) > n:
        s += f" … (+{len(allk)-n})"
    return s


def ctype(name):
    if " ABM - " in name:
        return "ABM (cible 1 entreprise)"
    if "Freemium" in name:
        return "Démo / inactif"
    return "Persona (cible un métier)"


wb = Workbook()

# ---------- Feuille 1 : Campagnes ----------
ws = wb.active
ws.title = "Campagnes"
cols = [
    ("Campagne", 34),
    ("Type", 22),
    ("Objectif (ce qu'on cherche)", 42),
    ("Cible implicite (d'après le nom)", 26),
    ("Deals créés", 11),
    ("dont REJETÉS (wrong_fit)", 12),
    ("% rejet", 9),
    ("Mots-clés de recherche actuels", 50),
    ("✏️ ACTION (Garder / Fusionner / Supprimer / Pause)", 26),
    ("✏️ CRITÈRES AFFINÉS — qui viser, qui exclure", 40),
    ("✏️ GÉO (régional / national)", 18),
    ("✏️ MOTS-CLÉS à AJOUTER ou RETIRER", 36),
]
for i, (title, w) in enumerate(cols, 1):
    c = ws.cell(1, i, title)
    c.fill = HEAD
    c.font = WHITEB
    c.alignment = WRAP
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 44
ws.freeze_panes = "A2"

rows = sorted(d["campaigns"], key=lambda x: -x["deals_total"])
r = 2
for c in rows:
    wrong = c["deals_by_outcome"].get("wrong_fit", 0)
    tot = c["deals_total"]
    pct = f"{100*wrong/tot:.0f}%" if tot else "—"
    cible = c["name"].split(" ABM - ")[-1] if " ABM - " in c["name"] else c["name"].replace("EKOALU - ", "")
    vals = [
        c["name"].replace("EKOALU - ", ""),
        ctype(c["name"]),
        c["objective"] or "(vide)",
        cible,
        tot,
        wrong,
        pct,
        kw_join(c["keywords"]) or "(aucun encore généré)",
        "", "", "", "",  # cases editables
    ]
    for i, v in enumerate(vals, 1):
        cell = ws.cell(r, i, v)
        cell.alignment = WRAP
        cell.border = THIN
        if i >= 9:
            cell.fill = EDIT
        elif i == 2:
            if "ABM" in vals[1]:
                cell.fill = ABMF
            elif "Persona" in vals[1]:
                cell.fill = PERS
    r += 1

# ---------- Feuille 2 : Historique mots-cles ----------
ws2 = wb.create_sheet("Historique mots-clés")
h2 = [("Date d'utilisation", 20), ("Campagne", 38), ("Mot-clé recherché", 50), ("Utilisé ?", 10)]
for i, (t, w) in enumerate(h2, 1):
    cc = ws2.cell(1, i, t)
    cc.fill = HEAD
    cc.font = WHITEB
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"
r = 2
for k in d["keywords_all"]:
    ws2.cell(r, 1, str(k["used_at"])[:19] if k["used_at"] else "(jamais)")
    ws2.cell(r, 2, k["campaign__name"].replace("EKOALU - ", ""))
    ws2.cell(r, 3, k["keyword"])
    ws2.cell(r, 4, "oui" if k["used"] else "non")
    r += 1

# ---------- Feuille 3 : Synthese ----------
ws3 = wb.create_sheet("À LIRE — Synthèse")
ws3.column_dimensions["A"].width = 110
lines = [
    ("LECTURE RAPIDE — diagnostic prospection EKOALU", True),
    ("", False),
    (f"Leads (personnes) en base : {d['leads_total']}", False),
    (f"Deals (= 1 test profil × campagne) : {d['deals_total']}", False),
    (f"  → dont REJETÉS 'wrong_fit' : {d['deals_by_outcome_global'].get('wrong_fit',0)} "
     f"({100*d['deals_by_outcome_global'].get('wrong_fit',0)/max(d['deals_total'],1):.1f}%)", False),
    (f"Campagnes : {len(d['campaigns'])} (dont ~42 ABM ciblant 1 entreprise précise)", False),
    (f"Mots-clés de recherche générés : {d['keywords_total']} (utilisés : {d['keywords_used']})", False),
    (f"Période de recherche : {str(d['keywords_used_first'])[:10]} → {str(d['keywords_used_last'])[:10]}", False),
    ("", False),
    ("LA CAUSE RACINE (en une phrase)", True),
    ("Chaque profil trouvé est testé contre PRESQUE TOUTES les campagnes. Comme 42 campagnes", False),
    ("ciblent chacune 1 seule entreprise (Léon Grosse, Bouygues, telle métallerie…), un bon", False),
    ("prospect 'colle' à 1 campagne et est rejeté par les 40 autres → ~97% de rejet MÉCANIQUE,", False),
    ("pas qualitatif. C'est ça qui gonfle le coût ET le risque LinkedIn EN MÊME TEMPS.", False),
    ("", False),
    ("CE QU'UN PROFIL NE GARDE PAS (vérifié en base)", True),
    ("Un profil ne mémorise PAS la campagne/recherche qui l'a trouvé. Le lien profil→campagne", False),
    ("n'existe qu'au moment du tri. C'est LE point à corriger pour pouvoir 'router' un profil", False),
    ("vers sa seule bonne campagne au lieu de le passer au crible de toutes.", False),
    ("", False),
    ("COMMENT REMPLIR L'ONGLET 'Campagnes'", True),
    ("Les colonnes JAUNES sont pour toi :", False),
    ("  • ACTION : Garder / Fusionner (avec quelle autre) / Supprimer / Pause", False),
    ("  • CRITÈRES AFFINÉS : qui viser précisément, qui exclure (en clair, pas technique)", False),
    ("  • GÉO : régional (Rhône-Alpes) ou national (si niche technique)", False),
    ("  • MOTS-CLÉS : termes de recherche à ajouter ou retirer", False),
    ("Je reprends ensuite tes réponses pour reconfigurer l'outil.", False),
]
for i, (txt, bold) in enumerate(lines, 1):
    cell = ws3.cell(i, 1, txt)
    cell.alignment = WRAP
    if bold:
        cell.font = Font(bold=True, size=13, color="1F4E78")

wb.save(OUT)
print("SAVED", OUT, "rows=", len(rows))
